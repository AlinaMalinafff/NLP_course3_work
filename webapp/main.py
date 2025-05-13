from flask import Flask, render_template, redirect, url_for, flash, request, send_from_directory, abort
from werkzeug.security import generate_password_hash, check_password_hash
from flask_sqlalchemy import SQLAlchemy
from flask_wtf import FlaskForm
from wtforms import StringField, PasswordField, SubmitField
from wtforms.validators import DataRequired, Length, Email, EqualTo
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from werkzeug.utils import secure_filename
import os
from datetime import datetime
from fpdf import FPDF
import uuid
import re
from docx import Document
import matplotlib.pyplot as plt
import io
import base64
import analyzer
from wtforms.validators import ValidationError
import re
from pymorphy2 import MorphAnalyzer

morph = MorphAnalyzer()
RUSSIAN_STOP_WORDS = {'и', 'в', 'во', 'не', 'что', 'он', 'на', 'я', 'с', 'со', 'как', 'а', 'то', 'все', 'она', 'так', 'его', 'но', 'да', 'ты', 'к', 'у', 'же', 'вы', 'за', 'бы', 'по', 'только', 'ее', 'мне', 'было', 'вот', 'от', 'меня', 'еще', 'нет', 'о', 'из', 'ему', 'теперь', 'когда', 'даже', 'ну', 'вдруг', 'ли', 'если', 'уже', 'или', 'ни', 'быть', 'был', 'него', 'до', 'вас', 'нибудь', 'опять', 'уж', 'вам', 'ведь', 'там', 'потом', 'себя', 'ничего', 'ей', 'может', 'они', 'тут', 'где', 'есть', 'надо', 'ней', 'для', 'мы', 'тебя', 'их', 'чем', 'была', 'сам', 'чтоб', 'без', 'будто', 'чего', 'раз', 'тоже', 'себе', 'под', 'будет', 'ж', 'тогда', 'кто', 'этот', 'того', 'потому', 'этого', 'какой', 'совсем', 'ним', 'здесь', 'этом', 'один', 'почти', 'мой', 'тем', 'чтобы', 'нее', 'сейчас', 'были', 'куда', 'зачем', 'всех', 'никогда', 'можно', 'при', 'наконец', 'два', 'об', 'другой', 'хоть', 'после', 'над', 'больше', 'тот', 'через', 'эти', 'нас', 'про', 'всего', 'них', 'какая', 'много', 'разве', 'три', 'эту', 'моя', 'впрочем', 'хорошо', 'свою', 'этой', 'перед', 'иногда', 'лучше', 'чуть', 'том', 'нельзя', 'такой', 'им', 'более', 'всегда', 'конечно', 'всю', 'между'}

def lemmatize_text(text):
    words = re.findall(r'\w+', text.lower())
    return [morph.parse(word)[0].normal_form for word in words if word not in RUSSIAN_STOP_WORDS]



app = Flask(__name__)

app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///site.db'
app.config['SECRET_KEY'] = 'my_very_secret_key_i_dont_know_why_i_need_it'
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['REPORTS_FOLDER'] = 'reports'
app.config['ALLOWED_EXTENSIONS'] = {'docx'}
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB максимум

db = SQLAlchemy(app)

# Создаем папки для загрузок и отчетов, если их нет
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
os.makedirs(app.config['REPORTS_FOLDER'], exist_ok=True)

# Проверка наличия шрифта Arial
if not os.path.exists('arial.ttf'):
    print("Ошибка: файл arial.ttf не найден в папке с main.py")
    print("Пожалуйста, поместите файл шрифта в ту же папку")
    exit(1)


# Модель пользователя
class User(db.Model, UserMixin):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(120), nullable=False)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password_hash = db.Column(db.String(60), nullable=False)

    def __repr__(self):
        return f"User('{self.username}', '{self.email}')"


# Модель загруженного файла
class UserFile(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    filename = db.Column(db.String(255), nullable=False)
    original_filename = db.Column(db.String(255), nullable=False)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    uploaded_at = db.Column(db.DateTime, default=db.func.current_timestamp())

    # Связь с отчетом
    report = db.relationship('FileReport', backref='file', uselist=False, lazy=True)
    user = db.relationship('User', backref=db.backref('files', lazy=True))


# Модель отчета
class FileReport(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    filename = db.Column(db.String(255), nullable=False)
    generated_at = db.Column(db.DateTime, default=db.func.current_timestamp())
    file_id = db.Column(db.Integer, db.ForeignKey('user_file.id'), nullable=False)


def validate_hse_email(form, field):
    """Валидатор для проверки, что email принадлежит домену hse.ru или edu.hse.ru"""
    email = field.data
    if not (email.endswith('@hse.ru') or email.endswith('@edu.hse.ru')):
        flash('Разрешены только email адреса доменов @hse.ru и @edu.hse.ru', 'danger')
        raise ValidationError('Разрешены только email адреса доменов @hse.ru и @edu.hse.ru')


# Формы
class RegistrationForm(FlaskForm):
    username = StringField('Имя пользователя', validators=[DataRequired(), Length(min=2, max=20)])
    email = StringField('Email', validators=[DataRequired(), Email(), validate_hse_email])
    password = PasswordField('Пароль', validators=[DataRequired()])
    confirm_password = PasswordField('Подтвердите пароль', validators=[DataRequired(), EqualTo('password')])
    submit = SubmitField('Зарегистрироваться')


class LoginForm(FlaskForm):
    email = StringField('Email', validators=[DataRequired(), Email(), validate_hse_email])
    password = PasswordField('Пароль', validators=[DataRequired()])
    submit = SubmitField('Войти')


# Инициализация Flask-Login
login_manager = LoginManager(app)
login_manager.login_view = 'login'


@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))


# Вспомогательные функции
def allowed_file(filename):
    return '.' in filename and \
        filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']


def parse_docx(filepath):
    """Парсинг DOCX файла и извлечение реплик респондента"""
    doc = Document(filepath)
    full_text = []
    for paragraph in doc.paragraphs:
        full_text.append(paragraph.text)
    full_text = "\n".join(full_text)

    # Ищем все реплики респондента
    respondent_pattern = r"Респондент\s*:\s*(.*?)(?=\s*(?:Интервьюер|$))"
    respondent_phrases = re.findall(respondent_pattern, full_text, flags=re.DOTALL)
    return [phrase.strip() for phrase in respondent_phrases]


def generate_report(filepath, original_filename):
    """Генерация PDF отчета с анализом"""
    # Парсинг и анализ текста
    paragraphs = parse_docx(filepath)
    analysis_results = []
    label_counts = {"Процесс": 0, "Результат": 0}
    model = analyzer.Model()

    for i, paragraph in enumerate(paragraphs, 1):
        print(paragraph)
        label = model.predict(paragraph)[0]
        print(type(label))
        print(label)
        label_counts[label] += 1
        analysis_results.append({
            "paragraph_num": i,
            "text": paragraph,
            "label": label
        })

    # Расчет статистики
    total = len(paragraphs) or 1
    percentages = {
        "Процесс": (label_counts["Процесс"] / total) * 100,
        "Результат": (label_counts["Результат"] / total) * 100
    }

    # Генерация графика
    plot_data = generate_plot(percentages)

    # Создаем PDF
    pdf = FPDF()
    pdf.add_page()

    # Настройка страницы
    pdf.set_margins(left=20, top=20, right=20)
    pdf.set_auto_page_break(auto=True, margin=15)

    # Шрифты
    pdf.add_font('ArialUnicode', '', 'arial.ttf', uni=True)
    pdf.add_font('ArialUnicode', 'B', 'arialbd.ttf', uni=True)
    pdf.add_font('ArialUnicode', 'I', 'ariali.ttf', uni=True)
    font_name = 'ArialUnicode'

    # Заголовок отчета
    pdf.set_font(font_name, 'B', 16)
    pdf.cell(0, 10, "Аналитический отчет", ln=1, align='C')
    pdf.set_font(font_name, '', 12)
    pdf.cell(0, 10, f"Файл: {original_filename}", ln=1, align='C')
    pdf.ln(15)

    # Статистика
    pdf.set_font(font_name, 'B', 14)
    pdf.cell(0, 10, "Общая статистика", ln=1)
    pdf.set_font(font_name, '', 12)

    # Добавляем график, если он есть
    if plot_data:
        try:
            # Создаем временный файл для графика
            temp_img_path = os.path.join(app.config['UPLOAD_FOLDER'], f"temp_plot_{uuid.uuid4().hex}.png")
            with open(temp_img_path, 'wb') as f:
                f.write(base64.b64decode(plot_data))

            # Вставляем изображение в PDF
            pdf.image(temp_img_path, x=50, w=110)
            pdf.ln(5)

            # Удаляем временный файл
            os.remove(temp_img_path)
        except Exception as e:
            print(f"Ошибка при вставке графика: {e}")

    pdf.cell(0, 8, f"Всего абзацев: {total}", ln=1)
    pdf.cell(0, 8, f"Процесс: {label_counts['Процесс']} ({percentages['Процесс']:.1f}%)", ln=1)
    pdf.cell(0, 8, f"Результат: {label_counts['Результат']} ({percentages['Результат']:.1f}%)", ln=1)
    pdf.ln(10)

    # Результаты по абзацам
    pdf.set_font(font_name, 'B', 14)
    pdf.cell(0, 10, "Детальный анализ по абзацам", ln=1)
    pdf.ln(5)

    for result in analysis_results:
        # Заголовок абзаца
        pdf.set_font(font_name, 'B', 12)
        pdf.set_fill_color(240, 240, 240)
        pdf.cell(0, 8, f"Абзац {result['paragraph_num']} ({result['label']})", ln=1, fill=True)

        # Текст абзаца
        pdf.set_font(font_name, '', 10)
        pdf.multi_cell(0, 6, result['text'])

        # Разделитель
        pdf.cell(0, 3, "", ln=1)
        pdf.set_draw_color(200, 200, 200)
        pdf.cell(0, 1, "", border='T', ln=1)
        pdf.ln(5)

    # Подвал
    pdf.set_y(-15)
    pdf.set_font(font_name, 'I', 8)
    pdf.cell(0, 10, f"Отчет сгенерирован {datetime.now().strftime('%Y-%m-%d %H:%M')}", 0, 0, 'C')

    # Сохраняем отчет
    report_filename = f"report_{uuid.uuid4().hex}.pdf"
    report_path = os.path.join(app.config['REPORTS_FOLDER'], report_filename)
    pdf.output(report_path)

    return report_filename


def generate_plot(percentages):
    """Генерация графика в base64"""
    plt.figure(figsize=(6, 4))
    plt.bar(percentages.keys(), percentages.values(), color=['blue', 'orange'])
    plt.title("Соотношение лейблов в документе")
    plt.ylabel("Процент")

    buf = io.BytesIO()
    plt.savefig(buf, format='png', dpi=100, bbox_inches='tight')
    buf.seek(0)
    plot_data = base64.b64encode(buf.getvalue()).decode('utf-8')
    plt.close()
    buf.close()
    return plot_data


# Маршруты
@app.route('/')
def home():
    return render_template('index.html')


@app.route('/info')
def info():
    return render_template('info.html')


@app.route('/login', methods=['GET', 'POST'])
def login():
    form = LoginForm()
    if form.validate_on_submit():
        email = form.email.data
        if not (email.endswith('@hse.ru') or email.endswith('@edu.hse.ru')):
            flash('Разрешены только email адреса доменов @hse.ru и @edu.hse.ru', 'danger')
            return redirect(url_for('register'))
        user = User.query.filter_by(email=form.email.data).first()
        if user and check_password_hash(user.password_hash, form.password.data):
            login_user(user)
            flash('Вы успешно вошли в систему!', 'success')
            return redirect(url_for('home_user'))
        else:
            flash('Неверный email или пароль', 'danger')
    return render_template('login.html', form=form)


@app.route('/home_user')
@login_required
def home_user():
    return render_template('home_user.html', user=current_user)


@app.route('/user_history')
@login_required
def user_history():
    user_files = UserFile.query.filter_by(user_id=current_user.id) \
        .order_by(UserFile.uploaded_at.desc()) \
        .all()
    return render_template('user_history.html', files=user_files, user=current_user)


@app.route('/analyzator', methods=['GET', 'POST'])
@login_required
def analyzator():
    if request.method == 'POST':
        if 'file' not in request.files:
            flash('Файл не был отправлен', 'danger')
            return redirect(url_for('analyzator'))

        file = request.files['file']

        if file.filename == '':
            flash('Не выбран файл для загрузки', 'danger')
            return redirect(url_for('analyzator'))

        if file and allowed_file(file.filename):
            unique_filename = f"{uuid.uuid4().hex}_{secure_filename(file.filename)}"
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], unique_filename)
            file.save(filepath)

            user_file = UserFile(
                filename=unique_filename,
                  original_filename=file.filename,
                user_id=current_user.id
            )
            db.session.add(user_file)
            db.session.commit()

            report_filename = generate_report(filepath, file.filename)

            report = FileReport(
                filename=report_filename,
                file_id=user_file.id
            )
            db.session.add(report)
            db.session.commit()

            flash('Файл успешно обработан и отчет сгенерирован!', 'success')
            return redirect(url_for('analyzator'))
        else:
            flash('Разрешены только файлы .docx', 'danger')

    user_files = UserFile.query.filter_by(user_id=current_user.id) \
        .order_by(UserFile.uploaded_at.desc()) \
        .all()
    return render_template('analyzator.html', user=current_user, files=user_files)


@app.route('/download_file/<int:file_id>')
@login_required
def download_file(file_id):
    user_file = UserFile.query.get_or_404(file_id)
    print("download")
    if user_file.user_id != current_user.id:
        print("first if")
        abort(403)
    print("before return")
    return send_from_directory(
        app.config['UPLOAD_FOLDER'],
        user_file.filename,
        as_attachment=True,
        download_name=user_file.original_filename
    )


@app.route('/download_report/<int:file_id>')
@login_required
def download_report(file_id):
    user_file = UserFile.query.get_or_404(file_id)
    if user_file.user_id != current_user.id or not user_file.report:
        abort(403)
    return send_from_directory(
        app.config['REPORTS_FOLDER'],
        user_file.report.filename,
        as_attachment=True,
        download_name=f"report_{user_file.original_filename}.pdf"
    )


@app.route('/logout')
@login_required
def logout():
    logout_user()
    flash('Вы вышли из системы', 'success')
    return redirect(url_for('login'))


@app.route('/register', methods=['GET', 'POST'])
def register():
    form = RegistrationForm()
    if form.validate_on_submit():
        email = form.email.data
        if not (email.endswith('@hse.ru') or email.endswith('@edu.hse.ru')):
            flash('Разрешены только email адреса доменов @hse.ru и @edu.hse.ru', 'danger')
            return redirect(url_for('register'))
        hashed_password = generate_password_hash(form.password.data, method='pbkdf2:sha256', salt_length=8)
        user = User(username=form.username.data, email=form.email.data, password_hash=hashed_password)
        db.session.add(user)
        db.session.commit()
        flash('Аккаунт успешно создан!', 'success')
        return redirect(url_for('login'))
    return render_template('register.html', form=form)


# Создание базы данных
with app.app_context():
    db.create_all()

if __name__ == '__main__':
    app.run(debug=True)
